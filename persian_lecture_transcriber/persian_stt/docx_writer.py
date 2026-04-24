"""Write a Word .docx with right-to-left Persian paragraphs."""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt


PERSIAN_FONT = "B Nazanin"         # install on Windows; alt: "Vazirmatn", "IRANSans"
PERSIAN_FONT_FALLBACK = "Arial"


def _set_rtl(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)


def _set_run_rtl(run) -> None:
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    rPr.append(rtl)


def write_rtl_docx(text: str, output_path: str, title: str | None = None) -> str:
    """Write `text` as an RTL Persian Word document. Blank lines separate paragraphs."""
    doc = Document()

    # Make the default style Persian-friendly.
    style = doc.styles["Normal"]
    style.font.name = PERSIAN_FONT
    style.font.size = Pt(14)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts")) or OxmlElement("w:rFonts")
    rfonts.set(qn("w:cs"), PERSIAN_FONT)
    rfonts.set(qn("w:ascii"), PERSIAN_FONT_FALLBACK)
    rfonts.set(qn("w:hAnsi"), PERSIAN_FONT_FALLBACK)
    if rfonts.getparent() is None:
        rpr.append(rfonts)

    if title:
        heading = doc.add_heading(level=1)
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        _set_rtl(heading)
        run = heading.add_run(title)
        _set_run_rtl(run)
        run.font.name = PERSIAN_FONT

    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        _set_rtl(p)
        run = p.add_run(block)
        _set_run_rtl(run)
        run.font.name = PERSIAN_FONT
        run.font.size = Pt(14)

    doc.save(output_path)
    return output_path
